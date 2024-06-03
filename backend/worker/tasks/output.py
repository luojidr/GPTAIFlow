# -*- coding: utf-8 -*-
# @Author: Bi Ying
# @Date:   2023-04-26 21:10:52
# @Last Modified by:   Bi Ying
# @Last Modified time: 2023-08-25 16:53:51
from pathlib import Path
from datetime import datetime

from config import IS_SECURITY_AUDITS
from utilities.security_audits import check_content
import yagmail
import openpyxl
import markdown2
from docx import Document
from docx.oxml.ns import qn
from htmldocx import HtmlToDocx

from utilities.workflow import Workflow
from utilities.print_utils import logger
from worker.tasks import task

from io import StringIO
import pandas as pd
import json
from models import WorkflowRunRecord
from utilities.print_utils import logger

@task
def text(
    workflow_data: dict,
    node_id: str,
):
    workflow = Workflow(workflow_data)
    text: str = workflow.get_node_field_value(node_id, "text")
    obj = WorkflowRunRecord.select().where(WorkflowRunRecord.rid==workflow.record_id).first()
    role = obj.user.role
    if eval(IS_SECURITY_AUDITS) and role == "BLACK_PERSON":
        msg = check_content(text)
        if msg:
            workflow.update_node_field_value(node_id, "text", msg)
    else:
        workflow.update_node_field_value(node_id, "text", text)
    return workflow.data


@task
def email(
    workflow_data: dict,
    node_id: str,
):
    workflow = Workflow(workflow_data)
    to_email: str = workflow.get_node_field_value(node_id, "to_email")
    subject: str = workflow.get_node_field_value(node_id, "subject")
    content_html: str = workflow.get_node_field_value(node_id, "content_html")
    obj = WorkflowRunRecord.select().where(WorkflowRunRecord.rid==workflow.record_id).first()
    role = obj.user.role
    if eval(IS_SECURITY_AUDITS) and role == "BLACK_PERSON":
        msg = check_content(content_html)
        if msg:
            raise ValueError(msg)
    yag = yagmail.SMTP(
        user=workflow.setting.get("email_user"),
        password=workflow.setting.get("email_password"),
        host=workflow.setting.get("email_smtp_host"),
        port=workflow.setting.get("email_smtp_port"),
        smtp_ssl=workflow.setting.get("email_smtp_ssl"),
    )
    email_send_result = yag.send(to_email, subject, [content_html])
    logger.info("email_send_result", email_send_result)
    return workflow.data


@task
def document(
    workflow_data: dict,
    node_id: str,
):
    workflow = Workflow(workflow_data)
    output_folder = Path(workflow.setting.get("output_folder"))
    file_name = workflow.get_node_field_value(node_id, "file_name")
    content = workflow.get_node_field_value(node_id, "content")
    obj = WorkflowRunRecord.select().where(WorkflowRunRecord.rid==workflow.record_id).first()
    role = obj.user.role
    if eval(IS_SECURITY_AUDITS) and role == "BLACK_PERSON":
        msg = check_content(content)
        if msg:
            contents = [msg]
        else:
            contents = [content]
    else:
        contents = [content]
    export_type = workflow.get_node_field_value(node_id, "export_type")

    local_file = output_folder / f"{file_name}{export_type}"
    if local_file.exists():
        local_file = output_folder / f"{file_name}_{datetime.now().strftime('%Y%m%d%H%M%S')}{export_type}"
    if export_type.endswith((".txt", ".md", ".html", ".json", ".csv")):
        with open(local_file, "w") as txt_file:
            txt_file.write("\n".join(contents))
    elif export_type.endswith(".docx"):
        content_str = "\n".join(contents)
        html_content = markdown2.markdown(content_str)
        new_parser = HtmlToDocx()
        document = Document()
        new_parser.add_html_to_document(html_content, document)
        for paragraph in document.paragraphs:
            for run in paragraph.runs:
                run.font.name = "微软雅黑"
                r = run._element.rPr.rFonts
                r.set(qn("w:eastAsia"), "微软雅黑")

        document.save(local_file)
    elif export_type.endswith(".xlsx"):
        content_str = "\n".join(contents)
        lines = content_str.split("\n")
        wb = openpyxl.Workbook()
        ws = wb.active
        for line in lines:
            ws.append(line.split(","))
        wb.save(local_file)

    file_full_path = str(local_file.resolve())
    workflow.update_node_field_value(node_id, "output", file_full_path)
    return workflow.data


# @task
# def audio(
#     workflow_data: dict,
#     node_id: str,
# ):
#     # TODO: Use local TTS like SpeechT5
#     workflow = Workflow(workflow_data)
#     content = workflow.get_node_field_value(node_id, "content")
#     output_type = workflow.get_node_field_value(node_id, "output_type")
#     # download_link = long_tts(content)
#     download_link = ""
#     workflow.update_node_field_value(node_id, "audio_url", download_link)
#     if output_type == "only_link":
#         workflow.update_node_field_value(node_id, "output", download_link)
#     elif output_type == "markdown":
#         workflow.update_node_field_value(node_id, "output", f"[{download_link}]({download_link})")
#     elif output_type == "html":
#         workflow.update_node_field_value(node_id, "output", f'<a href="{download_link}">{download_link}</a>')
#     return workflow.data


@task
def mindmap(
    workflow_data: dict,
    node_id: str,
):
    workflow = Workflow(workflow_data)
    content = workflow.get_node_field_value(node_id, "content")
    obj = WorkflowRunRecord.select().where(WorkflowRunRecord.rid==workflow.record_id).first()
    role = obj.user.role
    if eval(IS_SECURITY_AUDITS) and role == "BLACK_PERSON":
        msg = check_content(content)
        if msg:
            workflow.update_node_field_value(node_id, "content", msg)
    return workflow.data


@task
def mermaid(
    workflow_data: dict,
    node_id: str,
):
    workflow = Workflow(workflow_data)
    content = workflow.get_node_field_value(node_id, "content")
    obj = WorkflowRunRecord.select().where(WorkflowRunRecord.rid==workflow.record_id).first()
    role = obj.user.role
    if eval(IS_SECURITY_AUDITS) and role == "BLACK_PERSON":
        msg = check_content(content)
        if msg:
            workflow.update_node_field_value(node_id, "content", msg)
    return workflow.data


@task
def echarts(
    workflow_data: dict,
    node_id: str,
):
    workflow = Workflow(workflow_data)
    workflow.get_node_field_value(node_id, "option")
    return workflow.data


@task
def workflow_invoke_output(
    workflow_data: dict,
    node_id: str,
):
    workflow = Workflow(workflow_data)
    content = workflow.get_node_field_value(node_id, "value")
    obj = WorkflowRunRecord.select().where(WorkflowRunRecord.rid==workflow.record_id).first()
    role = obj.user.role
    if eval(IS_SECURITY_AUDITS) and role == "BLACK_PERSON":
        msg = check_content(content)
        if msg:
            workflow.update_node_field_value(node_id, "value", msg)
    return workflow.data

@task
def table(
    workflow_data: dict,
    node_id: str,
):
    workflow = Workflow(workflow_data)
    text: str = workflow.get_node_field_value(node_id, "text")
    obj = WorkflowRunRecord.select().where(WorkflowRunRecord.rid==workflow.record_id).first()
    role = obj.user.role
    if eval(IS_SECURITY_AUDITS) and role == "BLACK_PERSON":
        msg = check_content(text)
        if msg:
            text = msg
    if isinstance(text, list):
        df = pd.DataFrame(text)
    else:
        data_io = StringIO(text)
        df = pd.read_csv(data_io, delimiter='|', na_filter=False, keep_default_na=False)
    headers = df.columns.to_list()
    columns = [{'title':'Index',"dataIndex": 'Index',"showColumn": False, "resizable": True, "sorter": True, "sortDirections": ['descend', 'ascend']}]
    for i, header in enumerate(headers):
        if header == 'book_url':
            dataIndex = 'url'
        else:
            dataIndex = f'col_{i}'
        columns.append({'title':headers[i],"dataIndex": dataIndex,"showColumn": True, "resizable": True, "sorter": True, "sortDirections": ['descend', 'ascend']})
    df = df.rename(columns={column['title']:column['dataIndex'] for column in columns if column['title'] != 'Index'})
    row_list = [row._asdict() for row in df.itertuples()]
    
    final_dict = {"tableData":row_list, "tableColumns":columns}
    final_text = json.dumps(final_dict)

    workflow.update_node_field_value(node_id, "output", final_text)
    return workflow.data

@task
def audio(
    workflow_data: dict,
    node_id: str,
):
    workflow = Workflow(workflow_data)
    workflow.get_node_field_value(node_id, "byte_stream")
    return workflow.data